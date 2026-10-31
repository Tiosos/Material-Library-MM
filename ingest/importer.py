"""Supplier spreadsheets (CSV / XLSX) → UMR v2, imported according to a supplier mapping config.

Flow: read table → map headers → assemble drafts (material/codes/measures/prices/family) → validate →
(dry-run only reports) → apply writes to the DB (party / family / material / identifier / measure / price).

Design trade-offs:
* Idempotent: find an existing material by (identifier.kind, value, source_system) and update on a hit;
  only create when nothing is found, with material_uid / family_id generated from a stable hash, so
  repeated imports never create new records.
* Do not write material_revision / material_event: leave that to the DB triggers (single write path).
* Prices go into material_price separately; material.unit_cost is only the current cache.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from .mapping import Mapping, load_mapping, norm

CROCKFORD = "0123456789ABCDEFGHJKMNPQRSTVWXYZ"


def gen_uid(prefix: str, seed: str) -> str:
    """Stable ID: the same seed always yields the same uid (re-imports never create new records)."""
    n = int.from_bytes(hashlib.sha256(seed.encode("utf-8")).digest()[:16], "big")
    out = []
    for _ in range(26):
        n, r = divmod(n, 32)
        out.append(CROCKFORD[r])
    return f"{prefix}_{''.join(reversed(out))}"


def _slug(s: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(s or "").lower()).strip("-")


def _fmt(v: Any) -> str:
    """For template rendering: 110.0 → '110', None → ''."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ---------------------------------------------------------------------------
# Reading tables
# ---------------------------------------------------------------------------
def read_table(path: str | Path) -> tuple[list[str], list[dict]]:
    p = Path(path)
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        return _read_xlsx(p)
    return _read_csv(p)


def _read_csv(p: Path) -> tuple[list[str], list[dict]]:
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        rows = [dict(r) for r in reader]
        headers = list(reader.fieldnames or [])
    return headers, rows


def _read_xlsx(p: Path) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for r in it:
        if r is None or all(v is None for v in r):
            continue
        rows.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    wb.close()
    return headers, rows


# ---------------------------------------------------------------------------
# Drafts
# ---------------------------------------------------------------------------
@dataclass
class Draft:
    row_no: int
    part_no: str = ""
    name: str = ""
    category_code: str = ""
    unit: str | None = None
    currency: str | None = None
    unit_cost: float | None = None
    brand: str | None = None
    supplier: str | None = None
    family_key: str | None = None
    family_name: str | None = None
    family_axis: dict = field(default_factory=dict)
    attributes: dict = field(default_factory=dict)
    measures: list[dict] = field(default_factory=list)
    identifiers: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    matched_headers: dict = field(default_factory=dict)

    def as_dict(self) -> dict:
        return {
            "row": self.row_no,
            "part_no": self.part_no,
            "name": self.name,
            "category_code": self.category_code,
            "family_key": self.family_key,
            "family_name": self.family_name,
            "attributes": self.attributes,
            "measures": self.measures,
            "identifiers": self.identifiers,
            "unit_cost": self.unit_cost,
            "brand": self.brand,
            "supplier": self.supplier,
            "errors": self.errors,
            "warnings": self.warnings,
        }


def build_drafts(headers: list[str], rows: Iterable[dict], m: Mapping) -> tuple[list[Draft], dict]:
    resolved, unmatched = m.resolve(headers)
    missing = [f["key"] for f in m.fields if f.get("required") and f["key"] not in resolved]
    if missing:
        raise SystemExit(
            f"[{m.vendor}] required columns not matched: {missing}\n"
            f"  actual headers: {headers}\n"
            f"  please add them to aliases in mappings/{m.vendor}.yaml"
        )

    extra_mode = m.data.get("extra_headers", "ignore")
    defaults = m.defaults
    name_tpl = m.data.get("name_template", "{part_no}")
    family_cfg = m.data.get("family") or {}

    drafts: list[Draft] = []
    for i, row in enumerate(rows, start=2):
        d = Draft(row_no=i)
        values: dict[str, Any] = {}
        for key, spec in m.columns.items():
            values[key] = m.value(row, resolved.get(key), spec)

        d.matched_headers = {k: resolved[k] for k in resolved}

        # --- Identifiers and columns ---
        for key, spec in m.columns.items():
            v = values.get(key)
            if v in (None, ""):
                continue
            tgt = spec["target"]
            if tgt == "identifier":
                d.identifiers.append({
                    "kind": spec.get("kind", "supplier_code"),
                    "value": str(v),
                    "source_system": spec.get("source_system", m.source_system),
                })
                if spec.get("kind") in ("sku", "supplier_code"):
                    d.part_no = d.part_no or str(v)
            elif tgt == "attribute":
                d.attributes[key] = v
            elif tgt == "measure":
                num = m.value(row, resolved.get(key), {**spec, "transform": "number"})
                if num is None:
                    d.warnings.append(f"{key}: cannot convert to a number ({v!r})")
                else:
                    d.measures.append({
                        "key": spec.get("measure_key", key),
                        "value": num,
                        "unit": spec.get("unit"),
                        "source_system": m.source_system,
                        "source_ref": spec.get("source_ref"),
                    })
            elif tgt == "column":
                col = spec.get("column", key)
                if col == "unit_cost":
                    num = m.value(row, resolved.get(key), {**spec, "transform": "number"})
                    d.unit_cost = num
                else:
                    values[f"__col_{col}"] = v

        if extra_mode == "attributes":
            for h in unmatched:
                v = row.get(h)
                if v not in (None, ""):
                    d.attributes[f"src.{_slug(h)}"] = v

        # --- Assembly ---
        d.category_code = values.get("__col_category_code") or defaults.get("category_code", "generic")
        d.unit = values.get("__col_unit") or defaults.get("unit")
        d.currency = values.get("__col_currency") or defaults.get("currency", "AUD")
        d.brand = values.get("__col_brand") or defaults.get("brand")
        d.supplier = values.get("__col_supplier") or defaults.get("supplier")
        d.name = (values.get("__col_name") or name_tpl.format(**{
            k: _fmt(v) for k, v in {**defaults, **values}.items()
        })).strip()
        d.name = re.sub(r"\s+", " ", d.name)

        if not d.name:
            d.errors.append("name is empty")
        if not d.part_no:
            d.errors.append("no code column")

        # --- Product family ---
        if family_cfg:
            keys = family_cfg.get("key") or []
            parts = [_slug(values.get(k, "")) for k in keys]
            if not any(parts):
                d.warnings.append("all family keys are empty; skipping family grouping")
            else:
                fam_tpl = family_cfg.get("name_template", "{brand} " + " ".join(f"{{{k}}}" for k in keys))
                brand_name = d.brand or ""
                d.family_key = "|".join([_slug(d.brand or m.vendor), *parts])
                d.family_name = re.sub(r"\s+", " ", fam_tpl.format(
                    brand=brand_name, **{k: _fmt(values.get(k)) for k in keys})).strip()
                for ax in family_cfg.get("axis", []):
                    if values.get(ax) in (None, ""):
                        continue
                    # variant axis keys prefer measure_key (thickness → thickness_mm),
                    # so v_family lines up with the measure table
                    spec = m.columns.get(ax) or {}
                    axis_key = spec.get("measure_key", ax) if spec.get("target") == "measure" else ax
                    d.family_axis[axis_key] = values[ax]

        drafts.append(d)
    return drafts, {"resolved": resolved, "unmatched": unmatched}


# ---------------------------------------------------------------------------
# Applying
# ---------------------------------------------------------------------------
def apply(conn, drafts: list[Draft], m: Mapping) -> dict:
    stats = {"new": 0, "updated": 0, "skipped": 0, "measures": 0, "identifiers": 0, "families": 0}
    with conn.cursor() as cur:
        cur.execute("SELECT set_config('ml.actor', %s, true)", (f"import:{m.vendor}",))
        cur.execute("SELECT set_config('ml.source', %s, true)", (m.source_system,))
        cur.execute("SELECT code FROM material_category")
        valid_cats = {r[0] for r in cur.fetchall()}

        for d in drafts:
            if d.errors:
                stats["skipped"] += 1
                continue
            if d.category_code not in valid_cats:
                d.errors.append(f"unknown category {d.category_code!r}")
                stats["skipped"] += 1
                continue

            brand_id = _ensure_party(cur, d.brand, "brand", m.defaults.get("country")) if d.brand else None
            supplier_id = _ensure_party(cur, d.supplier, "supplier", m.defaults.get("country")) if d.supplier else None

            family_id = None
            if d.family_key:
                family_id = _ensure_family(cur, d, brand_id)
                stats["families"] += 1

            uid = _find_or_insert(cur, d, brand_id, supplier_id, family_id)
            if uid is None:
                stats["skipped"] += 1
                continue
            stats["new" if uid.new else "updated"] += 1

            for ident in d.identifiers:
                cur.execute(
                    """INSERT INTO material_identifier (material_uid, kind, value, source_system, is_primary)
                       VALUES (%s, %s, %s, %s, %s)
                       ON CONFLICT (kind, value, source_system) DO NOTHING""",
                    (uid.uid, ident["kind"], ident["value"], ident["source_system"], True),
                )
                stats["identifiers"] += 1

            for ms in d.measures:
                # keep only one preferred row per (material, measure) so repeated imports don't
                # trip the unique index
                cur.execute(
                    """UPDATE material_measure SET is_preferred = false
                        WHERE material_uid = %s AND key = %s AND is_preferred""",
                    (uid.uid, ms["key"]),
                )
                cur.execute(
                    """INSERT INTO material_measure
                         (material_uid, key, value_num, ucum, source_system, source_ref, is_preferred)
                       VALUES (%s, %s, %s, %s, %s, %s, true)
                       ON CONFLICT (material_uid, key, source_system, valid_from)
                       DO UPDATE SET value_num = EXCLUDED.value_num, ucum = EXCLUDED.ucum""",
                    (uid.uid, ms["key"], ms["value"], ms["unit"], ms["source_system"], ms.get("source_ref")),
                )
                stats["measures"] += 1

            if d.unit_cost is not None:
                cur.execute(
                    """INSERT INTO material_price (material_uid, supplier_id, unit_cost, currency, valid_from, source_system)
                       VALUES (%s, %s, %s, %s, %s, %s)
                       ON CONFLICT (material_uid, supplier_id, valid_from, qty_break, cost_basis) DO NOTHING""",
                    (uid.uid, supplier_id, d.unit_cost, d.currency or "AUD", date.today(), m.source_system),
                )
    return stats


@dataclass
class UidResult:
    uid: str
    new: bool


def _ensure_party(cur, name: str, kind: str, country: str | None) -> int:
    cur.execute("SELECT id FROM party WHERE lower(canonical_name) = lower(%s) AND kind IN (%s, 'both')",
                (name, kind))
    row = cur.fetchone()
    if row:
        return row[0]
    cur.execute("SELECT party_id FROM party_alias WHERE lower(alias) = lower(%s)", (name,))
    row = cur.fetchone()
    if row:
        return row[0]
    cur.execute("INSERT INTO party (kind, canonical_name, country) VALUES (%s, %s, %s) RETURNING id",
                (kind, name, country))
    return cur.fetchone()[0]


def _ensure_family(cur, d: Draft, brand_id: int | None) -> str:
    cur.execute("SELECT family_id FROM material_family WHERE family_key = %s", (d.family_key,))
    row = cur.fetchone()
    if row:
        return row[0]
    fid = gen_uid("fam", d.family_key)
    cur.execute(
        """INSERT INTO material_family (family_id, family_key, brand_id, display_name, axis_keys)
           VALUES (%s, %s, %s, %s, %s)
           ON CONFLICT (family_id) DO NOTHING""",
        (fid, d.family_key, brand_id, d.family_name or d.family_key, list(d.family_axis.keys())),
    )
    return fid


def _find_or_insert(cur, d: Draft, brand_id, supplier_id, family_id) -> UidResult | None:
    primary = d.identifiers[0] if d.identifiers else None
    if primary:
        cur.execute(
            "SELECT material_uid FROM material_identifier WHERE kind=%s AND value=%s AND source_system=%s",
            (primary["kind"], primary["value"], primary["source_system"]),
        )
        row = cur.fetchone()
        if row:
            cur.execute(
                """UPDATE material SET
                     name = %s, category_code = %s, brand_id = %s, default_supplier_id = %s,
                     family_id = %s, variant_values = %s, unit = %s, unit_cost = %s,
                     currency = %s, attributes = material.attributes || %s::jsonb
                   WHERE material_uid = %s""",
                (d.name, d.category_code, brand_id, supplier_id, family_id,
                 json.dumps(d.family_axis), d.unit, d.unit_cost, d.currency or "AUD",
                 json.dumps(d.attributes), row[0]),
            )
            return UidResult(row[0], False)

    uid = gen_uid("mat", f"{d.identifiers[0]['source_system']}:{d.identifiers[0]['value']}"
                         if d.identifiers else f"{d.name}:{d.row_no}")
    cur.execute(
        """INSERT INTO material
             (material_uid, category_code, name, brand_id, default_supplier_id, family_id,
              variant_values, unit, unit_cost, currency, attributes, source_system, source_id)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
           ON CONFLICT (material_uid) DO NOTHING""",
        (uid, d.category_code, d.name, brand_id, supplier_id, family_id,
         json.dumps(d.family_axis), d.unit, d.unit_cost, d.currency or "AUD",
         json.dumps(d.attributes), primary["source_system"] if primary else None,
         primary["value"] if primary else None),
    )
    return UidResult(uid, True)


def plan_report(drafts: list[Draft], meta: dict) -> dict:
    return {
        "rows": len(drafts),
        "ok": sum(1 for d in drafts if not d.errors),
        "errors": sum(1 for d in drafts if d.errors),
        "warnings": sum(len(d.warnings) for d in drafts),
        "resolved": meta.get("resolved", {}),
        "unmatched": meta.get("unmatched", []),
        "sample": [d.as_dict() for d in drafts[:3]],
        "error_rows": [{"row": d.row_no, "errors": d.errors} for d in drafts if d.errors][:20],
    }
